"""Lectura local y extraccion prudente de datos desde archivos de CV."""

from dataclasses import dataclass
from pathlib import Path
import zipfile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from lxml.etree import XMLSyntaxError
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from jobpulse_v15.analysis import RuleBasedExtractor


MAX_CV_FILE_BYTES = 15 * 1024 * 1024
MAX_PDF_PAGES = 100
SUPPORTED_CV_SUFFIXES = {".docx", ".pdf", ".txt"}


@dataclass(frozen=True)
class CandidateProfileSuggestion:
    category: str
    value_code: str
    value_text: str
    qualifier: str | None = None
    education_status: str | None = None
    numeric_value: float | None = None
    unit: str | None = None
    evidence: str = ""


@dataclass(frozen=True)
class CandidateProfileImport:
    source_name: str
    suggestions: tuple[CandidateProfileSuggestion, ...]
    warnings: tuple[str, ...]


def read_cv_text(path: str | Path) -> str:
    """Lee texto de PDF, DOCX o TXT local sin conservar el contenido."""
    cv_path = Path(path)
    suffix = cv_path.suffix.lower()
    if suffix not in SUPPORTED_CV_SUFFIXES:
        raise ValueError("Formato no soportado. Selecciona un archivo PDF, DOCX o TXT.")
    try:
        size = cv_path.stat().st_size
    except OSError as exc:
        raise ValueError(f"No se pudo leer el archivo seleccionado: {exc}") from exc
    if size == 0:
        raise ValueError("El archivo del CV esta vacio.")
    if size > MAX_CV_FILE_BYTES:
        raise ValueError("El CV supera el limite local de 15 MB.")

    if suffix == ".txt":
        try:
            text = cv_path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError(
                "El TXT no usa UTF-8. Guardalo en UTF-8 e intenta de nuevo."
            ) from exc
        except OSError as exc:
            raise ValueError(f"No se pudo leer el archivo seleccionado: {exc}") from exc
    elif suffix == ".docx":
        try:
            document = Document(cv_path)
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            table_text = [
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ]
            text = "\n".join((*paragraphs, *table_text))
        except (
            PackageNotFoundError,
            zipfile.BadZipFile,
            OSError,
            ValueError,
            KeyError,
            XMLSyntaxError,
        ) as exc:
            raise ValueError(f"No se pudo leer el DOCX: {exc}") from exc
    else:
        try:
            reader = PdfReader(cv_path, strict=False)
            if reader.is_encrypted:
                raise ValueError(
                    "Los PDF protegidos con contrasena no son compatibles."
                )
            if len(reader.pages) > MAX_PDF_PAGES:
                raise ValueError("El PDF supera el limite local de 100 paginas.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        except ValueError:
            raise
        except (PdfReadError, OSError, KeyError, NotImplementedError) as exc:
            raise ValueError(f"No se pudo leer el PDF: {exc}") from exc

    if not text.strip():
        if suffix == ".pdf":
            raise ValueError(
                "No se pudo extraer texto del PDF. Si es un escaneo, requiere OCR; "
                "JobPulse no ejecuta OCR automaticamente."
            )
        raise ValueError("El archivo no contiene texto que se pueda extraer.")
    return text


def extract_candidate_profile(path: str | Path) -> CandidateProfileImport:
    """Propone hechos detectables localmente; la UI debe solicitar confirmacion."""
    cv_path = Path(path)
    text = read_cv_text(cv_path)
    extraction = RuleBasedExtractor().extract(text)
    suggestions: dict[tuple[str, str], CandidateProfileSuggestion] = {}
    experience_values: set[tuple[float, str, str]] = set()

    for requirement in extraction.requirements:
        category = str(requirement["category"])
        if category not in {"education", "experience", "language", "license", "skill"}:
            continue
        code = requirement.get("value_code")
        if category == "license" and not code:
            code = "driver_license"
        if category == "experience":
            minimum = requirement.get("minimum")
            unit = requirement.get("unit")
            if minimum is None or unit != "years":
                continue
            experience_values.add((float(minimum), unit, str(requirement["evidence"])))
            continue
        if not code:
            continue
        suggestion = CandidateProfileSuggestion(
            category=category,
            value_code=str(code),
            value_text=str(requirement["value_text"]),
            qualifier=requirement.get("qualifier"),
            education_status=(
                requirement.get("education_status")
                if category == "education"
                else None
            ),
            evidence=str(requirement["evidence"]),
        )
        suggestions[(category, str(code).casefold())] = suggestion

    if len(experience_values) == 1:
        years, unit, evidence = next(iter(experience_values))
        suggestions[("experience", "years_experience")] = CandidateProfileSuggestion(
            category="experience",
            value_code="years_experience",
            value_text=f"{years:g} {unit} de experiencia",
            numeric_value=years,
            unit=unit,
            evidence=evidence,
        )

    warnings = [
        "La extraccion es local y preliminar. Revisa cada dato y selecciona "
        "explicitamente cuales guardar."
    ]
    if extraction.field_statuses.get("skill") == "unknown":
        warnings.append(
            "Se detecto contexto de competencias no reconocido; revisa el CV "
            "y agrega los datos faltantes manualmente."
        )
    if len(experience_values) > 1:
        warnings.append(
            "Se encontraron varias cantidades de experiencia; no se propuso "
            "un total para evitar confundir experiencia general y especifica."
        )
    if extraction.analysis_status == "partial":
        warnings.append(
            "Algunos campos no pudieron determinarse con certeza; "
            "la ausencia de una sugerencia no significa que el dato falte en el CV."
        )
    if not suggestions:
        warnings.append(
            "No se identificaron datos compatibles con el vocabulario actual; "
            "puedes registrar el perfil manualmente."
        )
    return CandidateProfileImport(
        source_name=cv_path.name,
        suggestions=tuple(suggestions[key] for key in sorted(suggestions)),
        warnings=tuple(warnings),
    )
