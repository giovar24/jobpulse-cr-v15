import sys
import tempfile
import unittest
from pathlib import Path
from pypdf.generic import (
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
)

from docx import Document
from pypdf import PdfWriter

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from jobpulse_v15.candidate_profile import (
    extract_candidate_profile,
    read_cv_text,
)


class TestCandidateProfileImport(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.directory = Path(self.temp_dir.name)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_extracts_candidate_facts_from_local_txt(self):
        path = self.directory / "candidate.txt"
        path.write_text(
            "Skills: Python and SQL.\n"
            "English - advanced.\n"
            "Bachelor's degree completed.\n"
            "Driver's license class B1.\n"
            "5 years of experience.\n",
            encoding="utf-8",
        )

        result = extract_candidate_profile(path)

        facts = {
            (item.category, item.value_code): item
            for item in result.suggestions
        }
        self.assertEqual(
            set(facts),
            {
                ("skill", "python"),
                ("skill", "sql"),
                ("language", "en"),
                ("education", "bachelor"),
                ("license", "B1"),
                ("experience", "years_experience"),
            },
        )
        self.assertEqual(facts[("education", "bachelor")].education_status, "completed")
        self.assertEqual(
            facts[("experience", "years_experience")].numeric_value, 5.0
        )
        self.assertTrue(all(item.evidence for item in result.suggestions))
        self.assertTrue(result.warnings)

    def test_extracts_text_from_docx_paragraphs_and_tables(self):
        path = self.directory / "candidate.docx"
        document = Document()
        document.add_paragraph("Technical skills: Python.")
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "SQL skills required."
        document.save(path)

        text = read_cv_text(path)
        result = extract_candidate_profile(path)

        self.assertIn("Technical skills: Python.", text)
        self.assertIn("SQL skills required.", text)
        self.assertEqual(
            {
                item.value_code
                for item in result.suggestions
                if item.category == "skill"
            },
            {"python", "sql"},
        )

    def test_scanned_or_empty_pdf_reports_ocr_limitation(self):
        path = self.directory / "scan.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        with path.open("wb") as file:
            writer.write(file)

        with self.assertRaisesRegex(ValueError, "OCR"):
            read_cv_text(path)

    def test_password_protected_pdf_is_rejected_explicitly(self):
        path = self.directory / "protected.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        writer.encrypt("local-test-password")
        with path.open("wb") as file:
            writer.write(file)

        with self.assertRaisesRegex(ValueError, "contrasena"):
            read_cv_text(path)

    def test_extracts_text_from_searchable_pdf(self):
        path = self.directory / "searchable.pdf"
        writer = PdfWriter()
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        resources = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {NameObject("/F1"): writer._add_object(font)}
                )
            }
        )
        content = DecodedStreamObject()
        content.set_data(b"BT /F1 12 Tf 72 720 Td (Technical skills: Python) Tj ET")
        page[NameObject("/Resources")] = resources
        page[NameObject("/Contents")] = writer._add_object(content)
        with path.open("wb") as file:
            writer.write(file)

        result = extract_candidate_profile(path)

        self.assertTrue(
            any(item.category == "skill" and item.value_code == "python"
                for item in result.suggestions)
        )

    def test_malformed_docx_is_reported_as_a_read_error(self):
        path = self.directory / "malformed.docx"
        path.write_bytes(b"not a valid office document")

        with self.assertRaisesRegex(ValueError, "DOCX"):
            read_cv_text(path)

    def test_rejects_unsupported_format_and_invalid_utf8(self):
        unsupported = self.directory / "candidate.rtf"
        unsupported.write_text("CV", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "Formato no soportado"):
            read_cv_text(unsupported)

        invalid_utf8 = self.directory / "candidate.txt"
        invalid_utf8.write_bytes(b"\xff\xfe\xfa")
        with self.assertRaisesRegex(ValueError, "UTF-8"):
            read_cv_text(invalid_utf8)

    def test_rejects_files_above_the_local_size_limit(self):
        path = self.directory / "large.txt"
        with path.open("wb") as file:
            file.truncate(15 * 1024 * 1024 + 1)

        with self.assertRaisesRegex(ValueError, "15 MB"):
            read_cv_text(path)

    def test_multiple_experience_values_are_not_collapsed_to_one_claim(self):
        path = self.directory / "candidate.txt"
        path.write_text(
            "3 years of relevant experience. 8 years of professional experience.",
            encoding="utf-8",
        )

        result = extract_candidate_profile(path)

        self.assertFalse(
            any(item.category == "experience" for item in result.suggestions)
        )
        self.assertTrue(
            any("varias cantidades de experiencia" in warning for warning in result.warnings)
        )


if __name__ == "__main__":
    unittest.main()
